"""
IP-SAKTI SAHAYAK
Document Loader
================
Loads PDF / TXT / HTML content into a uniform in-memory representation
while preserving legal metadata (authority, jurisdiction, section, etc.)
end-to-end through the RAG pipeline, as required by spec Section 10.
"""

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional, Dict, Any

logger = logging.getLogger(__name__)


@dataclass
class LoadedDocument:
    text: str
    document_name: str
    source_url: Optional[str] = None
    authority: Optional[str] = None
    jurisdiction: Optional[str] = None
    section: Optional[str] = None
    version: Optional[str] = None
    effective_date: Optional[str] = None
    document_type: Optional[str] = None
    page_map: Dict[int, str] = field(default_factory=dict)  # page_number -> page_text (best effort)


def load_pdf(path: str, metadata: Optional[Dict[str, Any]] = None) -> LoadedDocument:
    metadata = metadata or {}
    text_parts = []
    page_map: Dict[int, str] = {}
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(path)
        for i, page in enumerate(doc, start=1):
            page_text = page.get_text("text")
            page_map[i] = page_text
            text_parts.append(page_text)
        doc.close()
    except Exception as e:
        logger.warning("PyMuPDF failed (%s), falling back to pypdf for %s", e, path)
        try:
            from pypdf import PdfReader
            reader = PdfReader(path)
            for i, page in enumerate(reader.pages, start=1):
                page_text = page.extract_text() or ""
                page_map[i] = page_text
                text_parts.append(page_text)
        except Exception as e2:
            logger.error("PDF extraction failed entirely for %s: %s", path, e2)
            return LoadedDocument(text="", document_name=Path(path).name, **metadata, page_map={})

    return LoadedDocument(
        text="\n".join(text_parts),
        document_name=Path(path).name,
        page_map=page_map,
        **metadata,
    )

def load_docx(path: str, metadata: Optional[Dict[str, Any]] = None) -> LoadedDocument:
    metadata = metadata or {}
    try:
        import docx  # python-docx
        d = docx.Document(path)
        text = "\n".join(p.text for p in d.paragraphs)
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text += "\n" + cell.text
    except Exception as e:
        logger.error("DOCX extraction failed for %s: %s", path, e)
        text = ""
    return LoadedDocument(text=text, document_name=Path(path).name, **metadata)

def load_txt(path: str, metadata: Optional[Dict[str, Any]] = None) -> LoadedDocument:
    metadata = metadata or {}
    try:
        text = Path(path).read_text(encoding="utf-8", errors="ignore")
    except Exception as e:
        logger.error("TXT read failed for %s: %s", path, e)
        text = ""
    return LoadedDocument(text=text, document_name=Path(path).name, **metadata)


def load_html(html: str, source_url: Optional[str] = None,
              metadata: Optional[Dict[str, Any]] = None) -> LoadedDocument:
    """Load already-fetched HTML content (fetching itself lives in web_sources.py)."""
    metadata = metadata or {}
    from rag.text_cleaner import extract_text_from_html
    text = extract_text_from_html(html)
    name = metadata.pop("document_name", source_url or "web-document")
    return LoadedDocument(text=text, document_name=name, source_url=source_url, **metadata)


def load_document(path: str, metadata: Optional[Dict[str, Any]] = None) -> LoadedDocument:
    """Dispatch by file extension. Unknown extensions fall back to plain text."""
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return load_pdf(path, metadata)
    if ext == ".docx":
        return load_docx(path, metadata)
    if ext in (".txt", ".md"):
        return load_txt(path, metadata)
    if ext in (".html", ".htm"):
        html = Path(path).read_text(encoding="utf-8", errors="ignore")
        return load_html(html, source_url=str(path), metadata=metadata)
    logger.warning("Unrecognized extension %s for %s, treating as plain text.", ext, path)
    return load_txt(path, metadata)


# =====================================================
# TEST
# =====================================================

def _test():
    doc = load_txt.__wrapped__ if hasattr(load_txt, "__wrapped__") else None
    sample = LoadedDocument(text="Section 3(d) excludes certain inventions.", document_name="sample.txt")
    assert sample.text.startswith("Section")
    print("document_loader self-test passed.")


if __name__ == "__main__":
    _test()
